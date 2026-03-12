# sample_resume_generator_random.py

from docx import Document
from docx2pdf import convert
import os
import random

# Folder to save resumes
output_folder = "sample_resumes"
os.makedirs(output_folder, exist_ok=True)

# Role templates
roles = [
    "Data Scientist", "Data Analyst", "Python Developer", "Data Annotator",
    "Data Engineer", "Electrical Engineer", "Mechanical Engineer", "E&TC Engineer"
]

# Sample first and last names for randomization
first_names = ["Priya", "Rohan", "Ananya", "Sameer", "Kavya", "Arjun", "Mehul", "Nidhi", "Ritu", "Vikram"]
last_names = ["Sharma", "Mehta", "Gupta", "Khan", "Reddy", "Deshmukh", "Jain", "Kulkarni", "Sharma", "Singh"]

# Sample education for each role
education_templates = {
    "Data Scientist": ["M.Sc Data Science — University of Pune", "B.Tech Computer Science — IIT Bombay"],
    "Data Analyst": ["B.Sc Statistics — Delhi University"],
    "Python Developer": ["B.Tech Information Technology — VIT Pune"],
    "Data Annotator": ["B.A English — University of Hyderabad"],
    "Data Engineer": ["M.Tech Data Engineering — Anna University", "B.Tech IT — Amrita Vishwa Vidyapeetham"],
    "Electrical Engineer": ["B.E Electrical — College of Engineering Pune"],
    "Mechanical Engineer": ["B.Tech Mechanical — SVNIT Surat"],
    "E&TC Engineer": ["B.Tech E&TC — D Y Patil College, Pune"]
}

# Skills templates for each role
skills_templates = {
    "Data Scientist": ["Python", "R", "SQL", "Scikit-learn", "TensorFlow", "Pandas", "NumPy", "Matplotlib"],
    "Data Analyst": ["SQL", "Excel", "Tableau", "Power BI", "Python (Pandas, NumPy)"],
    "Python Developer": ["Python", "Flask", "Django", "REST APIs", "SQL/NoSQL", "Docker"],
    "Data Annotator": ["Labelbox", "CVAT", "Prodigy", "NLP Annotation", "Quality Assurance"],
    "Data Engineer": ["Python", "SQL", "Spark", "Hadoop", "Airflow", "AWS Redshift"],
    "Electrical Engineer": ["AutoCAD", "MATLAB", "ETAP", "Power Distribution", "PCB Design"],
    "Mechanical Engineer": ["SolidWorks", "CATIA", "FEA", "Thermodynamics", "Quality Control"],
    "E&TC Engineer": ["C/C++", "MATLAB", "PCB Design", "IoT", "Microcontrollers"]
}

# Function to generate random contact info
def random_email(name):
    domains = ["gmail.com", "yahoo.com", "outlook.com"]
    return f"{name.lower().replace(' ','')}@{random.choice(domains)}"

def random_phone():
    return f"+91-{random.randint(60000,99999)}-{random.randint(1000,9999)}"

# Function to create a DOCX resume
def create_resume(name, role):
    doc = Document()
    
    # Heading: Name
    doc.add_heading(name, level=0)
    
    # Role
    doc.add_paragraph(role, style='Intense Quote')
    
    # Contact info
    email = random_email(name)
    phone = random_phone()
    doc.add_paragraph(f"Email: {email} | Phone: {phone}")
    
    # Summary
    doc.add_heading("Summary", level=1)
    doc.add_paragraph(f"{role} with experience in relevant projects and technologies, eager to contribute effectively in a professional environment.")
    
    # Skills
    doc.add_heading("Skills", level=1)
    skills = skills_templates[role]
    doc.add_paragraph(", ".join(skills))
    
    # Experience
    doc.add_heading("Experience", level=1)
    doc.add_paragraph(f"{role} — Sample Company, City (2021–Present)")
    doc.add_paragraph("• Worked on projects relevant to the role, demonstrating skills and problem-solving.")
    doc.add_paragraph("• Collaborated with team members to deliver high-quality solutions.")
    
    # Education
    doc.add_heading("Education", level=1)
    for edu in education_templates[role]:
        doc.add_paragraph(edu)
    
    # Save DOCX
    filename = f"{name.replace(' ','_')}_{role.replace(' ','_')}"
    docx_path = os.path.join(output_folder, f"{filename}.docx")
    doc.save(docx_path)
    
    # Convert to PDF
    pdf_path = os.path.join(output_folder, f"{filename}.pdf")
    convert(docx_path, pdf_path)
    
    return docx_path, pdf_path

# Generate 10 random resumes
for i in range(10):
    first = random.choice(first_names)
    last = random.choice(last_names)
    role = random.choice(roles)
    name = f"{first} {last}"
    docx_file, pdf_file = create_resume(name, role)
    print(f"Created DOCX: {docx_file} | PDF: {pdf_file}")

print("All sample resumes generated successfully in 'sample_resumes' folder!")