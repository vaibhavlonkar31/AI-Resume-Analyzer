from pdfminer.high_level import extract_text
import docx

def extract_resume_text(file):

    text = ""

    if file.name.endswith(".pdf"):
        text = extract_text(file)

    elif file.name.endswith(".docx"):
        doc = docx.Document(file)
        text = " ".join([para.text for para in doc.paragraphs])

    return text

# import pdfplumber
# import docx

# def extract_resume_text(file):

#     text = ""

#     if file.name.endswith(".pdf"):
#         with pdfplumber.open(file) as pdf:
#             for page in pdf.pages:
#                 text += page.extract_text()

#     elif file.name.endswith(".docx"):
#         doc = docx.Document(file)
#         for para in doc.paragraphs:
#             text += para.text

#     return text